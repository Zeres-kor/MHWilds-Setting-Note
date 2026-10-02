from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
WEAPONS = ['공통', '대검', '태도', '한손검', '쌍검', '해머', '수렵피리', '랜스', '건랜스', '슬래시액스', '차지액스', '조충곤', '라이트보우건', '헤비보우건', '활']
GRADES = ['극종결', '종결', '준종결', '상급', '일반']
HEADERS = ['기준ID', '활성', '적용무기', '태그', '우선순위', '스킬판정', '스킬1', '레벨1', '스킬2', '레벨2', '스킬3', '레벨3', '슬롯판정', '무기슬롯1', '무기슬롯2', '무기슬롯3', '방어구슬롯1', '방어구슬롯2', '방어구슬롯3', '설명']
EXAMPLES = [
    ['COMMON-EX-01', '아니오', '공통', '준종결', 3, '이상', '혼신', 2, None, None, None, None, '이상', 0, 0, 0, 1, 1, 0, '형식 설명용. 추천·출현 가능성 검증 없음'],
    ['LS-EX-01', '아니오', '태도', '종결', 2, '정확', '납도술', 3, '혼신', 2, None, None, '정확', 0, 0, 0, 1, 1, 0, '스킬·슬롯 구성 정확 일치. 설명용, 출현 가능성 검증 없음'],
    ['GS-EX-01', '아니오', '대검', '준종결', 3, '이상', '공격', 3, None, None, None, None, '이상', 1, 0, 0, 1, 0, 0, '기준 이상 충족. 설명용, 출현 가능성 검증 없음'],
]

def make(path, examples=False):
    wb = Workbook()
    guide = wb.active
    guide.title = '사용법'
    lines = [
        ('항목', '설명'),
        ('양식 버전', '1.0 초안'),
        ('대상 시트', '태그기준 시트의 한 행이 하나의 규칙입니다.'),
        ('범위', '공통과 무기별 태그는 함께 적용됩니다.'),
        ('행의 의미', '한 행 안의 조건은 모두 충족. 같은 태그의 다른 행은 대안 조건.'),
        ('평가 태그', '극종결 > 종결 > 준종결 > 상급 > 일반. 미분류는 기준 불일치 상태입니다.'),
        ('우선순위', '작은 숫자가 먼저. 기본 순위는 극종결=1, 종결=2, 준종결=3, 상급=4, 일반=5.'),
        ('스킬판정', '이상: 추가 스킬 허용. 정확: 전체 스킬 구성 일치.'),
        ('슬롯판정', '이상/정확/무시. 무기는 방어구 슬롯을 대체하지 않습니다.'),
        ('슬롯 무시', '슬롯판정=무시이면 슬롯 칸은 모두 비워주세요.'),
        ('빈칸', '스킬과 레벨은 함께 입력하거나 함께 비워주세요. 슬롯 빈칸은 0입니다.'),
        ('예시 주의', '작성예시는 모두 비활성. 추천 기준 또는 게임 내 유효 호석 목록이 아닙니다.'),
        ('업로드', '전체 교체 또는 ID 병합 제안. 최종 앱에서 변경 미리보기 후 적용합니다.'),
        ('기본 판정', '극종결·종결은 스킬·슬롯 정확 일치, 준종결·상급·일반은 기준 이상 충족입니다.'),
        ('표시 정책', '범위별 최우선 태그 하나를 표시합니다. 범용과 무기별 태그는 함께 유지됩니다.'),
        ('데이터 입력', '수식 대신 값을 입력하세요. 기준ID와 헤더 이름을 유지하세요.'),
    ]
    for row in lines:
        guide.append(row)
    guide.column_dimensions['A'].width = 20
    guide.column_dimensions['B'].width = 100
    ws = wb.create_sheet('태그기준')
    ws.append(HEADERS)
    if examples:
        for row in EXAMPLES:
            ws.append(row)
    ws.freeze_panes = 'G2'
    ws.auto_filter.ref = f'A1:T{max(2, ws.max_row)}'
    for i, name in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = 16 if name != '설명' else 55
    for column, values in [('B', ['예','아니오']), ('C', WEAPONS), ('D', GRADES), ('F', ['이상','정확']), ('M', ['이상','정확','무시'])]:
        validation = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=False)
        validation.errorTitle = '입력값 확인'
        validation.error = '목록에 있는 값을 선택하세요.'
        validation.showErrorMessage = True
        ws.add_data_validation(validation)
        validation.add(f'{column}2:{column}10001')
    for column in ['E', 'H', 'J', 'L']:
        validation = DataValidation(type='whole', operator='greaterThanOrEqual', formula1=1, allow_blank=True)
        validation.showErrorMessage = True
        ws.add_data_validation(validation)
        validation.add(f'{column}2:{column}10001')
    for column in ['N','O','P','Q','R','S']:
        validation = DataValidation(type='whole', operator='between', formula1=0, formula2=3, allow_blank=True)
        validation.showErrorMessage = True
        ws.add_data_validation(validation)
        validation.add(f'{column}2:{column}10001')
    for sheet in wb:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='24445D')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    check = load_workbook(path)
    assert [c.value for c in check['태그기준'][1]] == HEADERS
    assert check['태그기준'].max_row == (4 if examples else 1)
    if examples:
        assert all(row[1] == '아니오' for row in check['태그기준'].iter_rows(min_row=2, values_only=True))
    check.close()
    print(path)

if __name__ == '__main__':
    make(ROOT / 'templates' / '태그기준_빈양식.xlsx')
    make(ROOT / 'templates' / '태그기준_작성예시.xlsx', examples=True)
